"""Parse uploaded .xlsx / .docx files into student rows.

Output shape per row (all raw strings, resolved to ids by the router):
    {"student_id", "name", "class_name", "section_name", "parent_name", "parent_phone"}
"""

from __future__ import annotations

import io
import re

_ALIASES = {
    "student_id": [
        "student id", "studentid", "id", "roll", "roll no", "rollno", "roll number",
        "admission no", "admission number", "adm no", "reg no", "gr no", "gr number",
        "enrollment no", "enrolment no",
    ],
    "name": [
        "name", "student name", "student", "full name", "pupil name", "child name",
        "name of student",
    ],
    "class_name": ["class", "grade", "standard", "std", "class name"],
    "section_name": ["section", "sec", "division", "div"],
    "parent_name": [
        "parent name", "parent", "guardian", "guardian name", "father name",
        "fathers name", "mother name", "parent guardian name", "parent or guardian",
    ],
    "parent_phone": [
        "parent phone", "phone", "mobile", "contact", "contact no", "contact number",
        "phone no", "phone number", "parent contact", "parent mobile", "parent number",
        "parent no", "parents number", "guardian phone", "guardian mobile",
    ],
}

_FIXED_ORDER = ["student_id", "name", "class_name", "section_name", "parent_name", "parent_phone"]


def _norm(text) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").strip().lower()).strip()


def _header_index(header_cells) -> dict[str, int]:
    norm = [_norm(h) for h in header_cells]
    idx: dict[str, int] = {}
    for field, aliases in _ALIASES.items():
        for alias in aliases:
            an = _norm(alias)
            if an in norm:
                idx[field] = norm.index(an)
                break
    return idx


def _row_to_student(cells, idx: dict[str, int]) -> dict:
    def get(field: str) -> str:
        j = idx.get(field)
        if j is None or j >= len(cells):
            return ""
        v = cells[j]
        return "" if v is None else str(v).strip()

    return {f: get(f) for f in _FIXED_ORDER}


def parse_students_excel(content: bytes) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = [
        list(r)
        for r in ws.iter_rows(values_only=True)
        if r and any(c is not None and str(c).strip() for c in r)
    ]
    if not rows:
        return []

    idx = _header_index(rows[0])
    if "name" not in idx:
        raise ValueError(
            "Could not find a 'Name' column in the spreadsheet. Expected a header "
            "row with columns like: student_id, name, class, section, parent_name, parent_phone."
        )

    out = []
    for r in rows[1:]:
        s = _row_to_student(r, idx)
        if any(s.values()):
            out.append(s)
    return out


def parse_students_docx(content: bytes) -> list[dict]:
    import docx

    doc = docx.Document(io.BytesIO(content))

    for table in doc.tables:
        if not table.rows:
            continue
        idx = _header_index([c.text for c in table.rows[0].cells])
        if "name" in idx and ("class_name" in idx or "student_id" in idx):
            out = []
            for row in table.rows[1:]:
                s = _row_to_student([c.text.strip() for c in row.cells], idx)
                if any(s.values()):
                    out.append(s)
            if out:
                return out

    # fallback: delimited text lines (tab / pipe / 2+ spaces)
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    rows = []
    for line in lines:
        parts = [p.strip() for p in re.split(r"\t|\s*\|\s*|\s{2,}", line) if p.strip()]
        if len(parts) >= 2:
            rows.append(parts)
    if not rows:
        return []

    idx = None
    if any(_norm(p) in {"name", "student name"} for p in rows[0]):
        idx = _header_index(rows[0])
        rows = rows[1:]

    out = []
    for parts in rows:
        if idx:
            s = _row_to_student(parts, idx)
        else:
            s = {k: (parts[i] if i < len(parts) else "") for i, k in enumerate(_FIXED_ORDER)}
        if s.get("name"):
            out.append(s)
    return out


def parse_students_upload(filename: str, content: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return parse_students_excel(content)
    if name.endswith(".docx"):
        return parse_students_docx(content)
    raise ValueError("Unsupported file type — upload a .docx or .xlsx file (.doc / .xls are not supported).")
