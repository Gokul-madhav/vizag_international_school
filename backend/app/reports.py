"""Admin analytics + examination results.

- `GET /api/admin/overview`  — dashboard numbers, graphs, recent activity
- `GET /api/exam/results`    — student assignment results, filterable by class / assignment
- `GET /api/exam/attempts/{id}` — one graded attempt with per-question detail
- `GET /api/exam/export`     — the filtered results as an .xlsx download
"""

from __future__ import annotations

import io
import os
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse

from .academics import repo as academic_repo
from .portal import _grade, attempt_repo
from .students import repo as students_repo

router = APIRouter(tags=["reports"])

# Base URL of the student-portal frontend, used to build a direct link to each
# test (`/student/a/{assignment_id}`). Override in .env if the frontend is
# hosted somewhere other than the Vite dev server.
FRONTEND_URL = os.getenv("FRONTEND_URL", "http://localhost:5173").strip().rstrip("/")


# --------------------------------------------------------------------------- #
# shared lookups
# --------------------------------------------------------------------------- #

def _context():
    tree = academic_repo.tree()
    students = students_repo.list()

    student_by_pk = {s["id"]: s for s in students}
    class_name = {c["id"]: c["name"] for c in tree["classes"]}
    section_name: dict[str, str] = {}
    for c in tree["classes"]:
        for sec in c["sections"]:
            section_name[sec["id"]] = sec["name"]

    assignment_meta: dict[str, dict] = {}
    for c in tree["classes"]:
        for subj in c["subjects"]:
            for topic in subj["topics"]:
                for a in topic["assignments"]:
                    assignment_meta[a["id"]] = {
                        "title": a["title"],
                        "subject": subj["name"],
                        "topic": topic["name"],
                        "class_id": c["id"],
                        "class_name": c["name"],
                        "question_count": a.get("question_count", 0),
                        "total_marks": a.get("total_marks", 0),
                        "max_attempts": a.get("max_attempts", 1),
                    }

    return {
        "tree": tree,
        "students": students,
        "student_by_pk": student_by_pk,
        "class_name": class_name,
        "section_name": section_name,
        "assignment_meta": assignment_meta,
    }


def _parse_dt(value: str):
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except Exception:  # noqa: BLE001
        return None


def _result_rows(ctx: dict, class_id: str | None = None, assignment_id: str | None = None):
    rows = []
    for at in attempt_repo.list_all():
        st = ctx["student_by_pk"].get(at["student_id"])
        am = ctx["assignment_meta"].get(at["assignment_id"])
        if not st or not am:
            continue
        if class_id and am["class_id"] != class_id:
            continue
        if assignment_id and at["assignment_id"] != assignment_id:
            continue
        rows.append(
            {
                "attempt_id": at.get("id"),
                "student_id": st.get("student_id") or "",
                "student_name": st["name"],
                "class": am["class_name"],
                "section": ctx["section_name"].get(st.get("section_id"), ""),
                "assignment": am["title"],
                "subject": am["subject"],
                "topic": am["topic"],
                "attempt_no": at["attempt_no"],
                "score": at["score"],
                "total": at["total"],
                "earned": at["earned"],
                "max_score": at["max_score"],
                "percentage": float(at["percentage"]),
                "submitted_at": at.get("submitted_at"),
            }
        )
    rows.sort(key=lambda r: (r["class"], r["assignment"], r["student_name"], r["attempt_no"]))
    return rows


def _assignment_rows(ctx: dict, class_id: str | None = None) -> list[dict]:
    """Every test across the school — name, chapter/subject, and a direct link
    (`/student/a/{id}`) a parent or student can be sent to open it."""
    rows = []
    for aid, meta in ctx["assignment_meta"].items():
        if class_id and meta["class_id"] != class_id:
            continue
        rows.append(
            {
                "id": aid,
                "class": meta["class_name"],
                "subject": meta["subject"],
                "topic": meta["topic"],
                "title": meta["title"],
                "question_count": meta["question_count"],
                "total_marks": meta["total_marks"],
                "max_attempts": meta["max_attempts"],
                "link": f"{FRONTEND_URL}/student/a/{aid}",
            }
        )
    rows.sort(key=lambda r: (r["class"].lower(), r["subject"].lower(), r["topic"].lower(), r["title"].lower()))
    return rows


# --------------------------------------------------------------------------- #
# admin dashboard overview
# --------------------------------------------------------------------------- #

@router.get("/api/admin/overview")
def admin_overview():
    ctx = _context()
    tree = ctx["tree"]
    attempts = attempt_repo.list_all()
    am = ctx["assignment_meta"]

    now = datetime.now(timezone.utc)
    day_keys = [(now.date() - timedelta(days=i)).isoformat() for i in range(29, -1, -1)]
    by_day = {d: 0 for d in day_keys}
    cutoff = now - timedelta(days=30)

    attempts_30d = 0
    pct_sum = 0.0
    pct_n = 0
    by_class: dict[str, list] = {}

    for at in attempts:
        sa = at.get("submitted_at") or ""
        if sa[:10] in by_day:
            by_day[sa[:10]] += 1
        dt = _parse_dt(sa)
        if dt and dt >= cutoff:
            attempts_30d += 1
        pct = float(at.get("percentage") or 0)
        pct_sum += pct
        pct_n += 1
        meta = am.get(at["assignment_id"])
        if meta:
            b = by_class.setdefault(meta["class_name"], [0, 0.0])
            b[0] += 1
            b[1] += pct

    students_by_class: dict[str, int] = {}
    for s in ctx["students"]:
        k = ctx["class_name"].get(s["class_id"], "—")
        students_by_class[k] = students_by_class.get(k, 0) + 1

    per_assignment: dict[str, list] = {}
    for at in attempts:
        per_assignment.setdefault(at["assignment_id"], []).append(float(at.get("percentage") or 0))
    top_assignments = []
    for aid, pcts in per_assignment.items():
        meta = am.get(aid)
        if not meta:
            continue
        top_assignments.append(
            {
                "title": meta["title"],
                "subject": meta["subject"],
                "class": meta["class_name"],
                "attempts": len(pcts),
                "avg": round(sum(pcts) / len(pcts), 1),
            }
        )
    top_assignments.sort(key=lambda r: -r["attempts"])

    recent = sorted(attempts, key=lambda a: a.get("submitted_at") or "", reverse=True)[:8]
    recent_results = []
    for at in recent:
        st = ctx["student_by_pk"].get(at["student_id"])
        meta = am.get(at["assignment_id"])
        if not st or not meta:
            continue
        recent_results.append(
            {
                "student": st["name"],
                "assignment": meta["title"],
                "percentage": float(at["percentage"]),
                "submitted_at": at.get("submitted_at"),
            }
        )

    return {
        "totals": {
            "students": len(ctx["students"]),
            "classes": len(tree["classes"]),
            "sections": sum(len(c["sections"]) for c in tree["classes"]),
            "subjects": sum(len(c["subjects"]) for c in tree["classes"]),
            "topics": sum(len(s["topics"]) for c in tree["classes"] for s in c["subjects"]),
            "assignments": len(am),
            "questions": sum(m["question_count"] for m in am.values()),
            "attempts": len(attempts),
            "attempts_30d": attempts_30d,
        },
        "avg_score": round(pct_sum / pct_n, 1) if pct_n else None,
        "attempts_by_day": [{"date": d, "count": by_day[d]} for d in day_keys],
        "attempts_by_class": [
            {"class": k, "count": v[0], "avg": round(v[1] / v[0], 1)}
            for k, v in sorted(by_class.items())
        ],
        "students_by_class": [
            {"class": k, "count": v} for k, v in sorted(students_by_class.items())
        ],
        "top_assignments": top_assignments[:5],
        "recent_results": recent_results,
    }


# --------------------------------------------------------------------------- #
# examination results
# --------------------------------------------------------------------------- #

@router.get("/api/exam/results")
def exam_results(class_id: str | None = None, assignment_id: str | None = None):
    ctx = _context()
    return {"results": _result_rows(ctx, class_id or None, assignment_id or None)}


@router.get("/api/exam/attempts/{attempt_id}")
def exam_attempt_detail(attempt_id: str):
    at = attempt_repo.get(attempt_id)
    if not at:
        raise HTTPException(status_code=404, detail="Attempt not found.")
    ctx = _context()
    st = ctx["student_by_pk"].get(at["student_id"])
    meta = ctx["assignment_meta"].get(at["assignment_id"])
    try:
        full = academic_repo.get_assignment(at["assignment_id"])
    except KeyError:
        raise HTTPException(status_code=404, detail="Assignment no longer exists.")
    graded = _grade(full["questions"], at.get("answers") or {})
    return {
        "student_name": st["name"] if st else "—",
        "student_id": (st.get("student_id") if st else "") or "",
        "class": meta["class_name"] if meta else "—",
        "assignment": meta["title"] if meta else full.get("title"),
        "subject": meta["subject"] if meta else full.get("subject"),
        "topic": meta["topic"] if meta else full.get("topic"),
        "attempt_no": at["attempt_no"],
        "submitted_at": at.get("submitted_at"),
        "score": graded["score"],
        "total": graded["total"],
        "earned": graded["earned"],
        "max_score": graded["max_score"],
        "percentage": graded["percentage"],
        "results": graded["results"],
    }


@router.get("/api/exam/export")
def exam_export(class_id: str | None = None, assignment_id: str | None = None):
    import openpyxl

    ctx = _context()
    rows = _result_rows(ctx, class_id or None, assignment_id or None)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    headers = [
        "Student ID", "Student Name", "Class", "Section", "Assignment", "Subject",
        "Topic", "Attempt #", "Correct", "Total Qs", "Marks Earned", "Max Marks",
        "Percentage", "Submitted At",
    ]
    ws.append(headers)
    for r in rows:
        ws.append(
            [
                r["student_id"], r["student_name"], r["class"], r["section"],
                r["assignment"], r["subject"], r["topic"], r["attempt_no"],
                r["score"], r["total"], r["earned"], r["max_score"],
                r["percentage"], (r["submitted_at"] or "")[:19].replace("T", " "),
            ]
        )
    for i, _ in enumerate(headers, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    label = "all-classes"
    if class_id:
        label = ctx["class_name"].get(class_id, "class").replace(" ", "_")
    filename = f"results_{label}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# --------------------------------------------------------------------------- #
# test / assignment catalog — names, chapters, direct links
# --------------------------------------------------------------------------- #

@router.get("/api/reports/tests")
def list_tests(class_id: str | None = None):
    ctx = _context()
    return {"tests": _assignment_rows(ctx, class_id or None)}


@router.get("/api/reports/tests/export")
def export_tests(class_id: str | None = None):
    import openpyxl

    ctx = _context()
    rows = _assignment_rows(ctx, class_id or None)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tests"
    headers = [
        "Class", "Subject", "Chapter / Topic", "Test Name",
        "Questions", "Total Marks", "Max Attempts", "Link",
    ]
    ws.append(headers)
    for r in rows:
        ws.append(
            [
                r["class"], r["subject"], r["topic"], r["title"],
                r["question_count"], r["total_marks"], r["max_attempts"], r["link"],
            ]
        )
    for i, _ in enumerate(headers, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    label = "all-classes"
    if class_id:
        label = ctx["class_name"].get(class_id, "class").replace(" ", "_")
    filename = f"tests_{label}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
